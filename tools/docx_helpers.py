"""python-docx helpers + a rough page-height estimator (no Word/LibreOffice available to paginate)."""
import math, re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Mm, RGBColor
from PIL import Image

FONT = "Times New Roman"
PAGE_H_PT = (297 / 25.4 - 1.5) * 72      # usable text height
CHARS_PER_LINE = int(__import__('os').environ.get('RELICAST_CPL', 84))                       # TNR 12 pt over 6.02 in
LINE_15 = 20.7                            # 12 pt TNR at 1.5 spacing


def set_font(style, name=FONT, size=None, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is not None:
        rpr.remove(rf)
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    rpr.insert(0, rf)
    if size: style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    if color is not None: style.font.color.rgb = color


def run_font(run, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)


def add_field(par, instr, text="1", dirty=False):
    def fc(t):
        r = par.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        if dirty and t == "begin": e.set(qn("w:dirty"), "true")
        r._r.append(e); run_font(r); return r
    fc("begin")
    r = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " %s " % instr; r._r.append(it)
    fc("separate")
    rr = par.add_run(text); run_font(rr)
    fc("end")


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)


def _insert(parent, el, before_tags):
    for t in before_tags:
        ref = parent.find(qn(t))
        if ref is not None:
            ref.addprevious(el); return
    parent.append(el)


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    _insert(tblPr, m, ["w:tblLook"])


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + k); e.set(qn("w:val"), "nil"); b.append(e)
    _insert(tblPr, b, ["w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"])


def page_border(section, val="double", sz=8, space=18):
    sp = section._sectPr
    pb = OxmlElement("w:pgBorders"); pb.set(qn("w:offsetFrom"), "text")
    for k in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + k); e.set(qn("w:val"), val); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), str(space)); e.set(qn("w:color"), "000000"); pb.append(e)
    # pgBorders must come after pgMar
    mar = sp.find(qn("w:pgMar")); mar.addnext(pb)


def pgnum(section, fmt="decimal", start=1):
    sp = section._sectPr
    for e in sp.findall(qn("w:pgNumType")): sp.remove(e)
    e = OxmlElement("w:pgNumType"); e.set(qn("w:fmt"), fmt); e.set(qn("w:start"), str(start))
    _insert(sp, e, ["w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid"])


def setup_section(sec):
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Inches(1.25), Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    sec.header_distance = sec.footer_distance = Inches(0.3)


_tok = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)")


def rich(par, text, size=None, font=None, bold=False, italic=False):
    for part in _tok.split(text):
        if not part: continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**"): part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: part, i = part[1:-1], True
        r = par.add_run(part); r.bold = b; r.italic = i
        run_font(r, font or FONT)
        if size: r.font.size = Pt(size)
    return par


class Est:
    """Accumulates approximate vertical extent (pt) per chapter and heading positions."""
    def __init__(self):
        self.ch = []            # list of dict(title, h)
        self.heads = []         # (level, text, chapter_index, offset_pt)
        self.cur = None

    def new_chapter(self, title):
        self.cur = {"title": title, "h": 0.0}; self.ch.append(self.cur)

    def add(self, h): self.cur["h"] += h

    def pages(self):
        return [max(1, math.ceil(c["h"] / PAGE_H_PT)) for c in self.ch]

    def head_page(self, ci, off):
        start = 1 + sum(self.pages()[:ci])
        return start + int(off // PAGE_H_PT)


def text_h(text, cpl=CHARS_PER_LINE, line=LINE_15, after=6):
    plain = re.sub(r"\*", "", text)
    return max(1, math.ceil(len(plain) / cpl)) * line + after


def img_ratio(path):
    with Image.open(path) as im:
        return im.height / im.width
