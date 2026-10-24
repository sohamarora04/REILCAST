"""Build ReliCast_Report.docx.   usage: make_report.py [experiments.json] [out.docx]"""
import os, sys, math, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *
from docx_helpers import *
from report_text import write_content


class B:
    """Builder: writes into a python-docx Document and tallies estimated height."""
    def __init__(self, doc, est, labels=None):
        self.doc, self.est, self.ch_no, self.fig_no, self.tab_no, self.eq_no = doc, est, 0, 0, 0, 0
        self.old_labels = labels or {}
        self.labels = {}

    def sub(self, text):
        return re.sub(r"\{\{(\w+)\}\}", lambda m: self.old_labels.get(m.group(1), "?"), text)

    # ---- structure
    def h1(self, title):
        self.ch_no += 1; self.fig_no = self.tab_no = 0
        self.est.new_chapter(title)
        p = self.doc.add_heading("%d. %s" % (self.ch_no, title.upper()), level=1)
        self.est.heads.append((1, "%d. %s" % (self.ch_no, title.upper()), len(self.est.ch) - 1, 0))
        self.est.add(60)
        return p

    def h2(self, title):
        n = "%d.%d" % (self.ch_no, sum(1 for h in self.est.heads if h[0] == 2 and h[2] == len(self.est.ch) - 1) + 1)
        self.doc.add_heading("%s  %s" % (n, title), level=2)
        self.est.heads.append((2, "%s  %s" % (n, title), len(self.est.ch) - 1, self.est.cur["h"]))
        self.est.add(LINE_15 + 8)

    def p(self, text, align="j", keep=False, after=6):
        par = self.doc.add_paragraph()
        par.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
        par.paragraph_format.space_after = Pt(after)
        if keep: par.paragraph_format.keep_with_next = True
        text = self.sub(text)
        rich(par, text)
        self.est.add(text_h(text, after=after))
        return par

    def bullets(self, items, style="List Bullet"):
        for t in items:
            par = self.doc.add_paragraph(style=style)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            rich(par, t)
            self.est.add(text_h(t, cpl=78, after=2))
        self.est.add(4)

    def numbered(self, items):
        for i, t in enumerate(items, 1):
            par = self.doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.left_indent = Inches(0.35); par.paragraph_format.first_line_indent = Inches(-0.3)
            par.paragraph_format.space_after = Pt(2)
            rich(par, "%d.  %s" % (i, t))
            self.est.add(text_h(t, cpl=78, after=2))
        self.est.add(4)

    def eq(self, text):
        self.eq_no += 1
        par = self.doc.add_paragraph()
        par.paragraph_format.tab_stops.add_tab_stop(Inches(3.0), alignment=1)
        par.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), alignment=2)
        par.paragraph_format.space_after = Pt(2); par.paragraph_format.line_spacing = 1.15
        par.add_run("\t")
        r = par.add_run(text); r.italic = True; run_font(r)
        par.add_run("\t(%d)" % self.eq_no)
        for r in par.runs: run_font(r)
        self.est.add(22 * (1 + len(text) // 70))

    def caption(self, text, kind, key=None):
        if kind == "Figure": self.fig_no += 1; n = self.fig_no
        elif kind == "Algorithm": self.alg_no = getattr(self, "alg_no", 0) + 1; n = self.alg_no
        else: self.tab_no += 1; n = self.tab_no
        if key: self.labels[key] = "%d.%d" % (self.ch_no, n)
        par = self.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0; par.paragraph_format.space_after = Pt(8); par.paragraph_format.space_before = Pt(2)
        r = par.add_run("%s %d.%d: " % (kind, self.ch_no, n)); r.bold = True; r.font.size = Pt(11); run_font(r)
        r2 = par.add_run(text); r2.font.size = Pt(11); run_font(r2)
        if kind in ("Table", "Algorithm"): par.paragraph_format.keep_with_next = True
        self.est.add(math.ceil((len(text) + 12) / 92) * 13 + 10)

    def figure(self, path, caption, width=5.6, key=None):
        if not os.path.exists(path):
            self.p("[Figure missing: %s]" % os.path.basename(path)); return
        r_ = img_ratio(path); maxh = 6.4
        w = min(width, maxh / r_)
        par = self.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1.0; par.paragraph_format.space_after = Pt(2); par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Inches(w))
        self.est.add(w * r_ * 72 + 6)
        self.caption(caption, "Figure", key)

    def placeholder(self, text, caption, key=None):
        t = self.doc.add_table(rows=1, cols=1); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; c.text = ""
        par = c.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.line_spacing = 1.0
        for _ in range(3): par.add_run("\n")
        rich(par, text, size=10, italic=True)
        for _ in range(3): par.add_run("\n")
        self.est.add(110)
        self.caption(caption, "Figure", key)

    def table(self, header, rows, caption, widths, size=10, align_num=False, key=None):
        self.caption(caption, "Table", key)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False; cell_margins(t)
        tot = 0
        for ri, row in enumerate([header] + rows):
            mx = 1
            for ci, txt in enumerate(row):
                c = t.rows[ri].cells[ci]; c.width = Inches(widths[ci]); c.text = ""
                par = c.paragraphs[0]; par.paragraph_format.line_spacing = 1.0; par.paragraph_format.space_after = Pt(0)
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT if (ci == 0 or not align_num) else WD_ALIGN_PARAGRAPH.CENTER
                rich(par, str(txt), size=size, bold=(ri == 0))
                if ri == 0: shade(c, "D9D9D9")
                mx = max(mx, math.ceil(len(str(txt)) / max(4, widths[ci] * 72 / (size * 0.47))))
            tot += mx * (size + 2.5) + 5
        self.est.add(tot + 10)
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4); sp.paragraph_format.line_spacing = 0.6
        self.est.add(8)

    def code(self, lines, caption=None, key=None):
        t = self.doc.add_table(rows=1, cols=1); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; cell_margins(t, 60, 60, 100, 100)
        c = t.rows[0].cells[0]; c.text = ""; shade(c, "F3F4F6")
        par = c.paragraphs[0]; par.paragraph_format.line_spacing = 1.0; par.paragraph_format.space_after = Pt(0)
        for i, ln in enumerate(lines):
            r = par.add_run(ln + ("\n" if i < len(lines) - 1 else "")); r.font.size = Pt(8.5); run_font(r, "Courier New")
        self.est.add(len(lines) * 10.4 + 16)
        if caption: self.caption(caption, "Algorithm", key)
        else:
            sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(2); self.est.add(6)


def make_styles(doc):
    st = doc.styles
    n = st["Normal"]; set_font(n, size=12)
    n.paragraph_format.line_spacing = 1.5; n.paragraph_format.space_after = Pt(6); n.paragraph_format.space_before = Pt(0)
    pf = n.paragraph_format; pf.widow_control = True
    h1 = st["Heading 1"]; set_font(h1, size=14, bold=True)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; h1.paragraph_format.page_break_before = True
    h1.paragraph_format.space_before = Pt(0); h1.paragraph_format.space_after = Pt(14); h1.paragraph_format.keep_with_next = True; h1.paragraph_format.line_spacing = 1.15
    h2 = st["Heading 2"]; set_font(h2, size=12, bold=True)
    h2.paragraph_format.space_before = Pt(8); h2.paragraph_format.space_after = Pt(4); h2.paragraph_format.keep_with_next = True; h2.paragraph_format.line_spacing = 1.5
    h3 = st["Heading 3"]; set_font(h3, size=12, bold=True, italic=True)
    for nm in ("List Bullet", "List Number"):
        set_font(st[nm], size=12); st[nm].paragraph_format.line_spacing = 1.5; st[nm].paragraph_format.space_after = Pt(2)
    for nm in ("Table Grid",):
        pass
    for nm in ("TOC 1", "TOC 2"):
        try: s = st[nm]
        except KeyError: s = st.add_style(nm, 1)
        s.base_style = st["Normal"]; set_font(s, size=12 if nm == "TOC 1" else 11, bold=(nm == "TOC 1"))
        s.paragraph_format.line_spacing = 1.15 if nm == "TOC 1" else 1.05
        s.paragraph_format.space_after = Pt(3 if nm == "TOC 1" else 0); s.paragraph_format.space_before = Pt(5 if nm == "TOC 1" else 0)
        s.paragraph_format.left_indent = Inches(0 if nm == "TOC 1" else 0.3)
        s.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), alignment=2, leader=1)


def fix_widths(t, widths):
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
        for c in t.columns[i].cells: c.width = Inches(w)


def centered(doc, text, size=12, bold=False, italic=False, after=6, before=0, spacing=1.15):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before); p.paragraph_format.line_spacing = spacing
    rich(p, text, size=size, bold=bold, italic=italic)
    return p


def cover(doc, T):
    logo = os.path.join(ASSETS, "logo.png")
    centered(doc, "B.M.S. COLLEGE OF ENGINEERING", 22, True, after=4, before=6)
    centered(doc, "(Autonomous Institute, Affiliated to VTU)", 15, True, after=2)
    centered(doc, "Bull Temple Road, Basavanagudi, Bengaluru - 560019", 15, True, after=22)
    p = centered(doc, "", after=22)
    if os.path.exists(logo): p.add_run().add_picture(logo, width=Inches(1.1))
    centered(doc, "Presentation Report under Alternate Assessment Tool (AAT)", 12, after=14)
    centered(doc, T["title"], 15, True, True, after=8)
    centered(doc, "Course: %s - %s" % (T["course_code"], T["course_title"]), 12, after=20)
    centered(doc, "Submitted in partial fulfilment of the requirements for the award of degree", 11, after=12)
    centered(doc, "BACHELOR OF ENGINEERING", 11, True, after=8)
    centered(doc, "IN", 11, True, after=8)
    centered(doc, "COMPUTER SCIENCE AND ENGINEERING", 11, True, after=10)
    centered(doc, "By", 11, after=6)
    t = doc.add_table(rows=len(T["members"]), cols=2); no_borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, m in enumerate(T["members"]):
        for j, k in enumerate(("name", "usn")):
            c = t.rows[i].cells[j]; c.width = Inches(2.6 if j == 0 else 1.8); c.text = ""
            pp = c.paragraphs[0]; pp.paragraph_format.line_spacing = 1.0; pp.paragraph_format.space_after = Pt(1)
            rich(pp, m[k], size=11)
    fix_widths(t, [2.8, 1.6])
    centered(doc, "Under the guidance of", 11, True, after=6, before=26)
    centered(doc, T["guide"], 12, True, after=1)
    centered(doc, T["guide_designation"], 12, after=18)
    centered(doc, T["department"], 14, True, after=2)
    centered(doc, T["year"], 14, True, after=6)


def certificate(doc, T):
    logo = os.path.join(ASSETS, "logo.png")
    t = doc.add_table(rows=1, cols=2); no_borders(t)
    fix_widths(t, [1.2, 4.8])
    a, b = t.rows[0].cells
    a.paragraphs[0].add_run().add_picture(logo, width=Inches(0.9))
    b.text = ""
    for i, (tx, sz, bd) in enumerate([("B.M.S. COLLEGE OF ENGINEERING", 14, True), ("( Autonomous Institute, Affiliated to VTU )", 12, True), ("Bull Temple Road, Basavanagudi,", 12, True), ("Bengaluru  – 560019", 12, True)]):
        pp = b.paragraphs[0] if i == 0 else b.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.space_after = Pt(0); pp.paragraph_format.line_spacing = 1.1
        rich(pp, tx, size=sz, bold=bd)
    centered(doc, T["department"], 14, True, after=34, before=18)
    p = centered(doc, "C E R T I F I C A T E", 14, True, after=28, before=10)
    names = ", ".join("**%s (%s)**" % (m["name"], m["usn"]) for m in T["members"])
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; par.paragraph_format.line_spacing = 1.5
    rich(par, "This is to certify that the Presentation Report under Alternate Assessment Tool entitled “***%s***” is a bona-fide work carried out by %s in partial fulfilment for the award of degree of Bachelor of Engineering in **Computer Science and Engineering** from **Visvesvaraya Technological University, Belagavi** during the year **%s**. It is certified that all corrections/suggestions indicated for Internal Assessments have been incorporated in the report deposited in the departmental library. The Presentation Report has been approved as it satisfies the academic requirements in respect of AAT prescribed for the Bachelor of Engineering Degree." % (T["title"], names, T["year"]))
    for _ in range(8): sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(6)
    t = doc.add_table(rows=2, cols=2); no_borders(t)
    fix_widths(t, [3.0, 3.0])
    for j, (n, r) in enumerate([(T["coordinator"], "Course Coordinator"), (T["hod"], "Professor and HOD")]):
        for i, tx in enumerate((n, r)):
            c = t.rows[i].cells[j]; c.text = ""; pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(0); pp.paragraph_format.line_spacing = 1.1
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            rich(pp, tx, size=11, bold=True)


def toc_page(doc, est):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10); p.paragraph_format.line_spacing = 1.15
    rich(p, "TABLE OF CONTENTS", size=14, bold=True)
    heads = est.heads if est else []
    # first/last paragraph carry the field begin/end so the whole list is one updatable TOC field
    n = len(heads)
    for i, (lvl, txt, ci, off) in enumerate(heads):
        par = doc.add_paragraph(style="TOC %d" % lvl)
        if i == 0:
            def fc(t, dirty=False):
                r = par.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
                if dirty: e.set(qn("w:dirty"), "true")
                r._r.append(e)
            fc("begin", True)
            r = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-2" \\h \\z \\u '; r._r.append(it)
            fc("separate")
        rich(par, "%s\t%d" % (txt, est.head_page(ci, off)), size=12 if lvl == 1 else 11, bold=(lvl == 1))
        if i == n - 1:
            r = par.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)


def footer_pagenum(sec):
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.0
    for r in list(p.runs): r._r.getparent().remove(r._r)
    add_field(p, "PAGE", "1")


def build(D, T, out):
    est1 = Est()
    d0 = Document(); make_styles(d0); b0 = B(d0, est1)
    write_content(b0, D, T)           # pass 1: measure
    est = Est()
    doc = Document(); make_styles(doc)
    s1 = doc.sections[0]; setup_section(s1)
    cover(doc, T)
    s2 = doc.add_section(WD_SECTION.NEW_PAGE); setup_section(s2)
    certificate(doc, T)
    s3 = doc.add_section(WD_SECTION.NEW_PAGE); setup_section(s3)
    toc_page(doc, est1)
    s4 = doc.add_section(WD_SECTION.NEW_PAGE); setup_section(s4)
    b = B(doc, est, b0.labels); write_content(b, D, T)
    # python-docx clones the running sectPr at every break, so configure each section explicitly afterwards
    secs = doc.sections
    for i, sx in enumerate(secs):
        sp = sx._sectPr
        for tag in ("w:pgBorders", "w:footerReference", "w:pgNumType"):
            for e in sp.findall(qn(tag)): sp.remove(e)
    page_border(secs[0])                       # boxed cover page only
    pgnum(secs[2], "lowerRoman", 1)            # TOC: i, ii ...
    pgnum(secs[3], "decimal", 1)               # chapters: 1, 2, ...
    footer_pagenum(secs[2])                    # page number in footer; section 4 inherits it
    # chapter 1 heading has page_break_before; avoid a blank page after the section break
    # remove the page_break_before on first heading (section break already starts a page)
    for par in doc.paragraphs:
        if par.style.name == "Heading 1":
            par.paragraph_format.page_break_before = False; break
    # ask Word to refresh the TOC on open
    st = doc.settings.element
    u = OxmlElement("w:updateFields"); u.set(qn("w:val"), "true")
    from docx_helpers import _insert
    _insert(st, u, ["w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "m:mathPr", "w:themeFontLang", "w:clrSchemeMapping", "w:shapeDefaults", "w:decimalSymbol", "w:listSeparator"])
    cp = doc.core_properties; cp.title = T["title"]; cp.author = ", ".join(m["name"] for m in T["members"]); cp.subject = "AAT Report - Network Programming"
    doc.save(out)
    return est1


def main():
    args = [a for a in sys.argv[1:]]
    path = args[0] if args else None
    out = args[1] if len(args) > 1 else os.path.join(ROOT, "ReliCast_Report.docx")
    D = Data(path); T = load_team()
    est = build(D, T, out)
    pg = est.pages()
    print("data:", D.path)
    for c, n in zip(est.ch, pg): print("  %-34s ~%d page(s)  (%.0f%% of last page)" % (c["title"], n, 100 * ((c["h"] % PAGE_H_PT) / PAGE_H_PT)))
    total = 2 + 1 + sum(pg)
    print("ESTIMATED total pages: cover+certificate+TOC (3) + chapters (%d) = %d" % (sum(pg), total))
    import docx
    d = docx.Document(out)
    print("reopened OK: %d paragraphs, %d tables, %d images" % (len(d.paragraphs), len(d.tables), len(d.inline_shapes)))
    print("wrote", out)


if __name__ == "__main__":
    main()
